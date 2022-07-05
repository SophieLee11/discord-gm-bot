from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference



def create_excel(headers, data):

	wb = Workbook()
	ws = wb.active

	ws.append(headers)

	for row in data:
		ws.append(row)

	width = len(data[0])
	height = len(data) + 1

	values = Reference(ws, min_col = 2, min_row = 1, max_col = width, max_row = height)
	chart = LineChart()

	name = 'Data'
	ws2 = wb.create_sheet(name)
	
	# ws2.add_chart(chart, 'A1')

	# chart.title = "Gm_data"
	# chart.y_axis.title = 'Value'
	# chart.x_axis.title = 'Date'
	# chart.x_axis.number_format = 'd-mmm'
	# chart.x_axis.majorTimeUnit = "days"

	# chart.add_data(values)

	wb.save("gm_data.xlsx")